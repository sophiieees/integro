# Genera la presentación de Íntegro para Oscar (Assignment 2) con python-docx.
# Uso: python herramientas/bitacora/generar_presentacion.py  (desde el venv con python-docx y Pillow)
# Salida: el .docx en el escritorio de Sophia. Los recortes temporales van a herramientas/bitacora/salida/ (ignorada por git).
import os
from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image

AQUI = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(AQUI, "..", "..")) + "/"
SCRATCH = os.path.join(AQUI, "salida") + "/"
os.makedirs(SCRATCH, exist_ok=True)
SALIDA = os.path.expanduser("~") + "/OneDrive/Escritorio/Integro_Presentacion_Oscar_2026-10-18.docx"
AZUL = RGBColor(0x0F, 0x3D, 0x52)
GRIS = RGBColor(0x5A, 0x68, 0x72)

# Recortes de las capturas para el documento
Image.open(REPO + "docs/diseno/capturas/portada-escritorio.png").crop((0, 0, 1366, 1330)).save(SCRATCH + "portada.png", optimize=True)
Image.open(REPO + "docs/diseno/capturas/ficha-escritorio.png").crop((0, 0, 1366, 1120)).save(SCRATCH + "ficha.png", optimize=True)

doc = Document()

# Página carta y márgenes
for s in doc.sections:
    s.page_width, s.page_height = Inches(8.5), Inches(11)
    s.left_margin = s.right_margin = Cm(2.3)
    s.top_margin = s.bottom_margin = Cm(2.0)

# Tipografía base y títulos en el azul de la marca
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
normal.paragraph_format.space_after = Pt(6)
for nombre, tam in (("Title", 26), ("Heading 1", 15), ("Heading 2", 12)):
    st = doc.styles[nombre]
    st.font.name = "Calibri"
    st.font.size = Pt(tam)
    st.font.color.rgb = AZUL
    st.font.bold = True
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
doc.styles["Title"].paragraph_format.space_after = Pt(2)
doc.styles["Heading 1"].paragraph_format.space_before = Pt(16)
doc.styles["Heading 2"].paragraph_format.space_before = Pt(10)


def parrafo(texto, negrita_inicio=None, estilo=None, cursiva=False, tam=None, color=None, alineacion=None):
    p = doc.add_paragraph(style=estilo)
    if negrita_inicio:
        r = p.add_run(negrita_inicio)
        r.bold = True
    r = p.add_run(texto)
    r.italic = cursiva
    if tam:
        r.font.size = Pt(tam)
    if color:
        r.font.color.rgb = color
    if alineacion is not None:
        p.alignment = alineacion
    return p


def vineta(texto, negrita_inicio=None):
    return parrafo(texto, negrita_inicio, estilo="List Bullet")


def sombrear(celda, hexcolor):
    tcPr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


# ---------- Encabezado del documento ----------
doc.add_paragraph("Íntegro", style="Title")
parrafo("Presentación del proyecto para el Assignment 2 de Applied Algorithmic Thinking", tam=13, color=AZUL)
parrafo("Sophia Pérez. 18 de octubre de 2026. Nombre provisional, sujeto a validación ante el IMPI.", tam=10, color=GRIS)

# ---------- 1 ----------
doc.add_heading("1. Por qué cambia el proyecto", level=1)
parrafo("Para el Assignment 1 usé LÍA, el producto de la empresa con la que colaboro. Para el Assignment 2 cambio a Íntegro por dos razones. "
        "La primera es confidencialidad: LÍA está en producción con datos de clientes reales y su código no puede vivir en un repositorio "
        "abierto a terceros ni modificarse para una tarea. La segunda es que Íntegro es un proyecto propio, en construcción desde el 5 de octubre "
        "de 2026, en el que soy responsable técnica: tengo un repositorio donde sí puedo construir una función real de principio a fin y "
        "documentar cada decisión.")
parrafo("El punto de partida es el mismo que pide el curso: una aplicación desplegada con un mockup público. La función del Assignment 2 "
        "se construye sobre ella, siguiendo los cuatro pasos de la tarea. Lo que ya existe se describe en la sección 3; lo que viene, en la 4.")

# ---------- 2 ----------
doc.add_heading("2. Qué es Íntegro", level=1)
parrafo("Íntegro es una plataforma de rentas de corta estancia para México, del tipo de Airbnb, con un modelo de ingresos distinto: "
        "el anfitrión paga una membresía mensual fija según el número de propiedades activas y no hay comisión por reserva. "
        "El huésped paga dentro de la plataforma y el dinero llega directo a la cuenta bancaria del anfitrión a través del procesador "
        "de pagos; la plataforma no custodia fondos. Lo que la distingue:")
vineta(" membresía mensual por propiedades activas (planes desde una propiedad hasta más de quince). Ningún porcentaje por reserva.", "Cero comisión:")
vineta(" el huésped ve noches, limpieza y total desde la búsqueda. No aparece un cargo de servicio al final.", "Precio total visible:")
vineta(" identidad con INE y rostro, cuenta CLABE y tarjeta a su nombre antes de publicar; documento de la propiedad y geolocalización por anuncio.", "Anfitrión verificado por capas:")
vineta(" los anfitriones nuevos cobran 24 horas después de la llegada del huésped; un reporte de fraude congela el depósito y el huésped recibe reembolso.", "Pago protegido:")
vineta(" el huésped entra con su teléfono por Bluetooth con una llave digital válida solo en sus fechas, con código numérico de respaldo. La apertura confirma la llegada real.", "Chapa inteligente:")
doc.add_heading("Equipo y estado", level=2)
parrafo("Somos tres socias. Valeria lleva negocio, legal y marca; Shakti, la arquitectura y la inteligencia artificial; yo, el diseño del MVP, "
        "la ciberseguridad y la construcción del backend y el frontend. El plan de construcción es de siete semanas, del 5 de octubre al "
        "20 de noviembre de 2026; hoy va la semana 2. La fase 1 supone 50 anfitriones, 150 propiedades y 500 reservas al mes en una sola ciudad, "
        "todavía por definir.")

# ---------- 3 ----------
doc.add_heading("3. Qué existe hoy", level=1)
parrafo("Un prototipo estático de la interfaz, en HTML, CSS y JavaScript sin dependencias, con cuatro páginas: la portada para el huésped "
        "(buscador sobre una imagen y seis propiedades de muestra), la ficha de una propiedad (precio total desglosado, verificaciones del "
        "anfitrión, zona aproximada enlazada a Google Maps y reseñas), la página \"Cómo funciona\" y la página para anfitriones (beneficios, "
        "planes y pasos para publicar). Los datos son ficticios y las fotos son marcadores que se reemplazan con solo guardar un archivo.")
parrafo("Una decisión deliberada: el prototipo no cotiza ni reserva. Muestra un total de ejemplo para dos noches. La primera función real "
        "(por ejemplo, la cotización por fechas) se construirá a partir de la especificación del Assignment 2, no antes, para que la tarea "
        "mida lo que debe medir.")
doc.add_picture(SCRATCH + "portada.png", width=Inches(6.0))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
parrafo("Portada: buscador sobre la imagen de bienvenida y propiedades con insignias de verificación y total sin cargos.", tam=9, cursiva=True, alineacion=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_picture(SCRATCH + "ficha.png", width=Inches(6.0))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
parrafo("Ficha de propiedad: capas de verificación del anfitrión y desglose con comisión de plataforma en cero.", tam=9, cursiva=True, alineacion=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_heading("Documentación en el repositorio", level=2)
vineta(" API primero, monolito modular, la plataforma no toca dinero ni documentos sensibles. Incluye módulos, modelo de datos, flujos, infraestructura y controles de seguridad.", "Arquitectura objetivo (docs/arquitectura.md):")
vineta(" once registros de decisiones de arquitectura, cada uno con contexto, decisión, consecuencias y fuente.", "ADR (docs/decisiones/):")
vineta(" paleta y tipografía en tokens, componentes, comparación con Airbnb y Wander, lista de pantallas y capturas.", "Sistema de diseño (docs/diseno/):")
vineta(" estructura de los cuatro pasos y 26 decisiones registradas hasta hoy (ver sección 4).", "Bitácora de prompts (PROMPT_LOG.md):")

# ---------- 4 ----------
doc.add_heading("4. Cómo se hará el Assignment 2 sobre Íntegro", level=1)
parrafo("Sigo los cuatro pasos del enunciado. El asistente de IA que uso es Claude Code dentro de VS Code.")
vineta(" en la Sesión 19 elijo una función entre las candidatas y la especifico en tres campos (qué hace, qué datos necesita, cómo se sabrá que funciona), auditada en pareja. Candidatas: cotización de una reserva con temporadas y estancia mínima; importación del calendario iCal de otra plataforma; detector de datos de contacto en mensajes (el fraude más común en estas plataformas); distancia entre la dirección declarada y el pin del anuncio con una API de geocodificación.", "Paso 1, especificar:")
vineta(" construyo la función citando mi especificación y registro cada propuesta de la IA que se salga de ella, con mi veredicto de aceptar, rechazar o revisar y la razón.", "Paso 2, construir:")
vineta(" corro al menos un caso de prueba real contra el deploy y registro el input exacto, el output copiado tal cual y mi veredicto con evidencia.", "Paso 3, validar:")
vineta(" entrego el enlace del deploy con la función integrada, el del repositorio y la bitácora completa, con la sección \"Uso de IA\".", "Paso 4, entregar:")
parrafo("La bitácora ya tiene evidencia del proceso que la rúbrica pide: diez propuestas de la IA pendientes de mi veredicto (por ejemplo, "
        "prototipo estático en lugar de Next.js, paleta de colores, omitir funciones de segunda fase) y dieciséis decisiones con mis veredictos "
        "citados, tomadas en cuatro revisiones de diseño el 18 de octubre (qué va en la portada, separar la página del anfitrión, mover secciones, "
        "cambiar colores, quitar separadores). El historial de commits explica cada decisión en el mensaje.")

doc.add_heading("Dónde está cada pieza", level=2)
filas = [
    ("Pieza del assignment", "Dónde está"),
    ("Funciones candidatas (Pre-Class Work S19)", "Sección 1 de PROMPT_LOG.md, al elegir"),
    ("Especificación de 3 campos, auditada (Checkpoint A)", "PROMPT_LOG.md, sección 1"),
    ("Registro aceptar / rechazar / revisar (Checkpoint B)", "PROMPT_LOG.md, sección 2"),
    ("Registro de validación con caso de prueba real", "PROMPT_LOG.md, sección 3"),
    ("Enlaces de deploy y repositorio", "PROMPT_LOG.md, sección 4"),
    ("Por qué la función resuelve el problema", "PROMPT_LOG.md, sección 5"),
    ("Declaración \"Uso de IA\"", "PROMPT_LOG.md, sección 6"),
    ("Código de la función", "Ruta anotada en la sección 1 al construirla"),
    ("Commits con decisiones", "Historial del repositorio (git log)"),
]
tabla = doc.add_table(rows=len(filas), cols=2)
tabla.style = "Table Grid"
tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
anchos = (Inches(3.4), Inches(3.0))
for i, (a, b) in enumerate(filas):
    c0, c1 = tabla.rows[i].cells
    c0.width, c1.width = anchos
    c0.text, c1.text = a, b
    for c in (c0, c1):
        for p in c.paragraphs:
            p.paragraph_format.space_after = Pt(2)
            for r in p.runs:
                r.font.size = Pt(10)
                if i == 0:
                    r.bold = True
                    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    if i == 0:
        sombrear(c0, "0F3D52")
        sombrear(c1, "0F3D52")
    elif i % 2 == 0:
        sombrear(c0, "F3F6F8")
        sombrear(c1, "F3F6F8")

# ---------- 5 ----------
doc.add_heading("5. Enlaces y acceso", level=1)
vineta(" https://github.com/sophiieees/integro", "Repositorio:")
vineta(" https://github.com/sophiieees/integro/blob/main/PROMPT_LOG.md", "Bitácora de prompts:")
vineta(" https://github.com/sophiieees/integro/blob/main/README.md (incluye la tabla \"Para el profesor\")", "README:")
vineta(" [enlace del deploy en Vercel, por agregar]", "Aplicación desplegada:")
parrafo("El repositorio es privado por confidencialidad del proyecto. Con su usuario de GitHub lo agrego como colaborador para que pueda "
        "revisar el código, la bitácora y el historial de commits.")

# ---------- 6 ----------
doc.add_heading("6. Uso de IA", level=1)
parrafo("Uso Claude Code como asistente de programación, como el curso lo permite. La IA propone; yo decido, pruebo y registro. El prototipo, "
        "la documentación técnica y este documento se hicieron con su ayuda. Las decisiones de producto y de diseño (qué muestra la portada, "
        "separar la página del anfitrión, qué funciones quedan fuera, colores, textos) son mías y están citadas en la bitácora con mis palabras.")

# ---------- 7. Anexo: bitácora de prompts (misma fuente que PROMPT_LOG.md) ----------
import sys
sys.path.insert(0, AQUI)
from bitacora_datos import PROMPTS, NOTA_TRANSPARENCIA  # noqa: E402


def vineta2(texto, negrita_inicio=None):
    return parrafo(texto, negrita_inicio, estilo="List Bullet 2")


doc.add_page_break()
doc.add_heading("7. Anexo: bitácora de prompts", level=1)
parrafo(NOTA_TRANSPARENCIA + " La versión viva está en PROMPT_LOG.md del repositorio; las secciones de especificación, "
        "validación y entrega final se llenan conforme avanza el assignment. Las decisiones marcadas como pendientes esperan mi veredicto.")

for p in PROMPTS:
    doc.add_heading("Prompt %d: %s (%s)" % (p["n"], p["titulo"], p["fecha"]), level=2)
    parrafo(p["rol"], "Rol. ")
    parrafo(p["tarea"], "Tarea. ")
    parrafo(p["contexto"], "Contexto. ")
    parrafo(p["salida"], "Salida esperada. ")
    r = p["respuesta"]
    parrafo("", "Respuesta de la IA")
    vineta(" %s. Commits: %s." % (r["fecha"], r["commits"]), "Fecha:")
    vineta("", "Carpetas y archivos, y por qué:")
    for ruta, porque in r["archivos"]:
        vineta2(" " + porque, ruta + ":")
    vineta("", "Sugerencias de la IA y por qué:")
    for sug, porque in r["sugerencias"]:
        vineta2(" Por qué: " + porque, sug)
    vineta(" " + r["lineas"], "Líneas o secciones modificadas:")
    if p["decisiones"]:
        parrafo("", "Decisiones de Sophia")
        filas = [("N", "Propuesta de la IA", "Veredicto", "Razón (en palabras de Sophia)")]
        for n, prop, porque, ver, razon, cambio in p["decisiones"]:
            filas.append((n, prop, ver, razon if razon else "(por escribir)"))
        t = doc.add_table(rows=len(filas), cols=4)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        anchos = (Inches(0.5), Inches(2.7), Inches(0.9), Inches(2.3))
        for i, fila in enumerate(filas):
            for j, texto in enumerate(fila):
                c = t.rows[i].cells[j]
                c.width = anchos[j]
                c.text = texto
                for par in c.paragraphs:
                    par.paragraph_format.space_after = Pt(1)
                    for run in par.runs:
                        run.font.size = Pt(9)
                        if i == 0:
                            run.bold = True
                            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                if i == 0:
                    sombrear(c, "0F3D52")
                elif i % 2 == 0:
                    sombrear(c, "F3F6F8")
        doc.add_paragraph()

doc.save(SALIDA)
print("guardado:", SALIDA)
